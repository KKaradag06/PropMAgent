import os
import re
from datetime import datetime
from typing import Optional, Dict, Any, List
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

REPORTS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "reports")

def _build_confirmation_docx(
    doc: Document,
    confirmation_data: Dict[str, Any],
    sender_name: str,
    sender_street: str,
    sender_city: str,
    is_private: bool
) -> None:
    """
    Erstellt die Elemente für Seite 2 im DOCX-Dokument:
    Offizielle Zustimmungserklärung / Bestätigungsschreiben für den Mieter zum Unterschreiben & Zurücksenden.
    """
    doc.add_page_break()

    let_date = confirmation_data.get("letter_date") or datetime.now().strftime("%d.%m.%Y")
    prop_name = confirmation_data.get("property_name") or ""
    unit_no = confirmation_data.get("unit_number") or ""
    tenant_names = confirmation_data.get("tenant_names") or []
    names_str = ", ".join(tenant_names) if tenant_names else "Mieter/in"
    conf_type = confirmation_data.get("type", "allgemein")

    # 1. Rücksende-Kopfzeile
    tbl_hdr = doc.add_table(rows=1, cols=2)
    tbl_hdr.autofit = False
    tbl_hdr.columns[0].width = Cm(9.5)
    tbl_hdr.columns[1].width = Cm(7.0)

    c_left = tbl_hdr.cell(0, 0)
    c_right = tbl_hdr.cell(0, 1)

    p_l = c_left.paragraphs[0]
    p_l.paragraph_format.space_after = Pt(2)
    p_l.paragraph_format.line_spacing = 1.0
    r_l1 = p_l.add_run("Rückantwort an den Absender:\n")
    r_l1.bold = True
    r_l1.font.name = "Arial"
    r_l1.font.size = Pt(8.5)
    r_l1.font.color.rgb = RGBColor(71, 85, 105)

    ret_suffix = " (Vermieter)\n" if is_private else "\n"
    r_l2 = p_l.add_run(f"{sender_name}{ret_suffix}{sender_street}\n{sender_city}")
    r_l2.font.name = "Arial"
    r_l2.font.size = Pt(9.0)
    r_l2.font.color.rgb = RGBColor(15, 23, 42)

    p_r = c_right.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_r.paragraph_format.space_after = Pt(2)
    p_r.paragraph_format.line_spacing = 1.0
    meta_lines = [f"Datum: {let_date}"]
    if prop_name:
        meta_lines.append(f"Mietobjekt: {prop_name}")
    if unit_no:
        meta_lines.append(f"Einheit: {unit_no}")
    r_r = p_r.add_run("\n".join(meta_lines))
    r_r.font.name = "Arial"
    r_r.font.size = Pt(9.0)
    r_r.font.color.rgb = RGBColor(51, 65, 85)

    # 2. Überschrift
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(14)
    p_title.paragraph_format.space_after = Pt(4)
    if conf_type == "mieterhoehung":
        title_text = "Zustimmungserklärung zur Mieterhöhung"
    elif conf_type == "nebenkosten":
        title_text = "Bestätigung der Anpassung der Nebenkostenvorauszahlung"
    else:
        title_text = "Bestätigungsschreiben / Rückantwort"
    r_t = p_title.add_run(title_text)
    r_t.bold = True
    r_t.font.name = "Arial"
    r_t.font.size = Pt(13.0)
    r_t.font.color.rgb = RGBColor(15, 23, 42)

    # 3. Referenzzeile
    p_ref = doc.add_paragraph()
    p_ref.paragraph_format.space_before = Pt(0)
    p_ref.paragraph_format.space_after = Pt(8)
    r_ref = p_ref.add_run(f"Mietobjekt: {prop_name}, {unit_no}  ·  Mieter: {names_str}  ·  Bezug: Schreiben vom {let_date}")
    r_ref.font.name = "Arial"
    r_ref.font.size = Pt(9.5)
    r_ref.font.color.rgb = RGBColor(71, 85, 105)

    # 4. Bestätigungstext und Beträge
    if conf_type == "mieterhoehung":
        eff_date = confirmation_data.get("effective_date", "")
        old_c = confirmation_data.get("old_cold", 0.0)
        diff_c = confirmation_data.get("diff_cold", 0.0)
        new_c = confirmation_data.get("new_cold", 0.0)
        old_nk = confirmation_data.get("old_nk", 0.0)
        new_tot = confirmation_data.get("new_total", 0.0)
        old_tot = confirmation_data.get("old_total", old_c + old_nk)
        deadline = confirmation_data.get("deadline_date", "")

        p_t1 = doc.add_paragraph()
        p_t1.paragraph_format.space_before = Pt(0)
        p_t1.paragraph_format.space_after = Pt(8)
        p_t1.paragraph_format.line_spacing = 1.0
        r_t1 = p_t1.add_run(
            f"Hiermit erkläre ich / erklären wir mein / unser uneingeschränktes Einverständnis mit der im Schreiben "
            f"vom {let_date} dargelegten Erhöhung der monatlichen Nettokaltmiete gemäß § 558 BGB."
        )
        r_t1.font.name = "Arial"
        r_t1.font.size = Pt(10.5)

        rows_data = [
            ["Mietbestandteil", "Bisher", "Erhöhung", f"Neu ab {eff_date}"],
            ["Nettokaltmiete", f"{old_c:.2f} €", f"+ {diff_c:.2f} €", f"{new_c:.2f} €"],
            ["Betriebskostenvorauszahlung", f"{old_nk:.2f} €", "unverändert", f"{old_nk:.2f} €"],
            ["Monatliche Gesamtmiete", f"{old_tot:.2f} €", f"+ {diff_c:.2f} €", f"{new_tot:.2f} €"]
        ]
        tbl = doc.add_table(rows=len(rows_data), cols=4)
        tbl.autofit = False
        widths = [Cm(6.0), Cm(3.2), Cm(3.3), Cm(4.0)]
        for r_i, r_data in enumerate(rows_data):
            row = tbl.rows[r_i]
            for c_i, val in enumerate(r_data):
                row.cells[c_i].width = widths[c_i]
                p_c = row.cells[c_i].paragraphs[0]
                p_c.paragraph_format.space_before = Pt(3)
                p_c.paragraph_format.space_after = Pt(3)
                r_c = p_c.add_run(val)
                r_c.font.name = "Arial"
                r_c.font.size = Pt(9.5)
                if r_i == 0 or r_i == len(rows_data) - 1 or c_i == 3:
                    r_c.bold = True

        p_t2 = doc.add_paragraph()
        p_t2.paragraph_format.space_before = Pt(8)
        p_t2.paragraph_format.space_after = Pt(6)
        p_t2.paragraph_format.line_spacing = 1.0
        r_t2 = p_t2.add_run(
            f"Ich / Wir verpflichten uns, die neue monatliche Gesamtmiete in Höhe von {new_tot:.2f} € "
            f"ab dem {eff_date} pünktlich zum Fälligkeitstermin auf das angegebene Mietkonto "
            f"zu entrichten und den bestehenden Dauerauftrag rechtzeitig anzupassen."
        )
        r_t2.font.name = "Arial"
        r_t2.font.size = Pt(10.5)

        dl_note = f" bis spätestens zum {deadline}" if deadline else ""
        p_note = doc.add_paragraph()
        p_note.paragraph_format.space_before = Pt(0)
        p_note.paragraph_format.space_after = Pt(8)
        r_note = p_note.add_run(
            f"Hinweis: Bitte senden Sie diese Erklärung{dl_note} ausgefüllt und von allen im "
            f"Mietvertrag genannten Personen unterzeichnet an den Vermieter zurück."
        )
        r_note.font.name = "Arial"
        r_note.font.size = Pt(9.0)
        r_note.font.italic = True
        r_note.font.color.rgb = RGBColor(100, 116, 139)

    elif conf_type == "nebenkosten":
        eff_date = confirmation_data.get("effective_date", "")
        old_c = confirmation_data.get("old_cold", 0.0)
        old_nk = confirmation_data.get("old_nk", 0.0)
        diff_nk = confirmation_data.get("diff_nk", 0.0)
        new_nk = confirmation_data.get("new_nk", 0.0)
        new_tot = confirmation_data.get("new_total", 0.0)
        old_tot = confirmation_data.get("old_total", old_c + old_nk)
        is_inc = confirmation_data.get("is_nk_inc", False)
        diff_sym = f"+ {diff_nk:.2f} €" if is_inc else f"- {diff_nk:.2f} €"

        p_t1 = doc.add_paragraph()
        p_t1.paragraph_format.space_before = Pt(0)
        p_t1.paragraph_format.space_after = Pt(8)
        p_t1.paragraph_format.line_spacing = 1.0
        r_t1 = p_t1.add_run(
            f"Hiermit bestätige ich / bestätigen wir den Erhalt Ihres Schreibens vom {let_date} und erklären "
            f"das Einverständnis mit der Anpassung der monatlichen Betriebskostenvorauszahlung gemäß § 560 Abs. 4 BGB."
        )
        r_t1.font.name = "Arial"
        r_t1.font.size = Pt(10.5)

        rows_data = [
            ["Mietbestandteil", "Bisher", "Anpassung", f"Neu ab {eff_date}"],
            ["Nettokaltmiete", f"{old_c:.2f} €", "unverändert", f"{old_c:.2f} €"],
            ["Betriebskostenvorauszahlung", f"{old_nk:.2f} €", diff_sym, f"{new_nk:.2f} €"],
            ["Monatliche Gesamtmiete", f"{old_tot:.2f} €", diff_sym, f"{new_tot:.2f} €"]
        ]
        tbl = doc.add_table(rows=len(rows_data), cols=4)
        tbl.autofit = False
        widths = [Cm(6.0), Cm(3.2), Cm(3.3), Cm(4.0)]
        for r_i, r_data in enumerate(rows_data):
            row = tbl.rows[r_i]
            for c_i, val in enumerate(r_data):
                row.cells[c_i].width = widths[c_i]
                p_c = row.cells[c_i].paragraphs[0]
                p_c.paragraph_format.space_before = Pt(3)
                p_c.paragraph_format.space_after = Pt(3)
                r_c = p_c.add_run(val)
                r_c.font.name = "Arial"
                r_c.font.size = Pt(9.5)
                if r_i == 0 or r_i == len(rows_data) - 1 or c_i == 3:
                    r_c.bold = True

        p_t2 = doc.add_paragraph()
        p_t2.paragraph_format.space_before = Pt(8)
        p_t2.paragraph_format.space_after = Pt(6)
        p_t2.paragraph_format.line_spacing = 1.0
        r_t2 = p_t2.add_run(
            f"Ich / Wir werden die monatlichen Mietzahlungen bzw. den bestehenden Dauerauftrag ab dem "
            f"{eff_date} rechtzeitig auf den neuen Gesamtbetrag von {new_tot:.2f} € anpassen."
        )
        r_t2.font.name = "Arial"
        r_t2.font.size = Pt(10.5)

        p_note = doc.add_paragraph()
        p_note.paragraph_format.space_before = Pt(0)
        p_note.paragraph_format.space_after = Pt(8)
        r_note = p_note.add_run(
            "Hinweis: Bitte senden Sie diese Bestätigung ausgefüllt und von allen Mietvertragsparteien unterzeichnet an den Vermieter zurück."
        )
        r_note.font.name = "Arial"
        r_note.font.size = Pt(9.0)
        r_note.font.italic = True
        r_note.font.color.rgb = RGBColor(100, 116, 139)

    else:
        subj = confirmation_data.get("subject", "")
        p_t1 = doc.add_paragraph()
        p_t1.paragraph_format.space_before = Pt(0)
        p_t1.paragraph_format.space_after = Pt(8)
        p_t1.paragraph_format.line_spacing = 1.0
        r_t1 = p_t1.add_run(
            f"Hiermit bestätige ich / bestätigen wir den ordnungsgemäßen Erhalt und die Kenntnisnahme "
            f"Ihres Schreibens vom {let_date}" + (f" betreffend {subj}." if subj else ".")
        )
        r_t1.font.name = "Arial"
        r_t1.font.size = Pt(10.5)

        p_t2 = doc.add_paragraph()
        p_t2.paragraph_format.space_before = Pt(0)
        p_t2.paragraph_format.space_after = Pt(8)
        p_t2.paragraph_format.line_spacing = 1.0
        r_t2 = p_t2.add_run(
            "Die im genannten Schreiben aufgeführten Mitteilungen bzw. Vereinbarungen wurden zur Kenntnis "
            "genommen und werden hiermit vollinhaltlich bestätigt."
        )
        r_t2.font.name = "Arial"
        r_t2.font.size = Pt(10.5)

        p_note = doc.add_paragraph()
        p_note.paragraph_format.space_before = Pt(0)
        p_note.paragraph_format.space_after = Pt(8)
        r_note = p_note.add_run(
            "Hinweis: Bitte senden Sie diese Empfangsbestätigung ausgefüllt und unterzeichnet an den Absender zurück."
        )
        r_note.font.name = "Arial"
        r_note.font.size = Pt(9.0)
        r_note.font.italic = True
        r_note.font.color.rgb = RGBColor(100, 116, 139)

    # 5. Unterschriften
    p_date = doc.add_paragraph()
    p_date.paragraph_format.space_before = Pt(14)
    p_date.paragraph_format.space_after = Pt(24)
    r_date = p_date.add_run("Ort, Datum: ____________________________________________________")
    r_date.font.name = "Arial"
    r_date.font.size = Pt(10.5)

    if len(tenant_names) <= 1:
        sig_name = tenant_names[0] if tenant_names else "Mieter/in"
        p_sig = doc.add_paragraph()
        p_sig.paragraph_format.space_before = Pt(0)
        p_sig.paragraph_format.space_after = Pt(0)
        r_s1 = p_sig.add_run("____________________________________________________\n")
        r_s1.font.name = "Arial"
        r_s1.font.size = Pt(10.5)
        r_s2 = p_sig.add_run(f"Unterschrift: {sig_name} (Mieter/in)")
        r_s2.font.name = "Arial"
        r_s2.font.size = Pt(10.5)
        r_s2.bold = True
    else:
        sig_tbl = doc.add_table(rows=1, cols=2)
        sig_tbl.autofit = False
        sig_tbl.columns[0].width = Cm(8.0)
        sig_tbl.columns[1].width = Cm(8.0)
        for idx, name in enumerate(tenant_names[:2]):
            cell = sig_tbl.cell(0, idx)
            p_cs = cell.paragraphs[0]
            p_cs.paragraph_format.space_before = Pt(0)
            p_cs.paragraph_format.space_after = Pt(0)
            r_cs1 = p_cs.add_run("________________________________________\n")
            r_cs1.font.name = "Arial"
            r_cs1.font.size = Pt(10.5)
            r_cs2 = p_cs.add_run(f"Unterschrift: {name}")
            r_cs2.font.name = "Arial"
            r_cs2.font.size = Pt(10.5)
            r_cs2.bold = True

def generate_formal_letter_docx(
    recipient_name: str,
    subject: str,
    letter_text: str,
    address_line: str = "",
    sender_name: Optional[str] = None,
    property_name: Optional[str] = None,
    sender_street: Optional[str] = None,
    sender_city: Optional[str] = None,
    sender_phone: Optional[str] = None,
    sender_email: Optional[str] = None,
    sign_person: Optional[str] = None,
    confirmation_data: Optional[Dict[str, Any]] = None
) -> str:
    """
    Erstellt ein formelles DOCX-Word-Dokument nach DIN 5008:2020 Form B.
    Nutzt standardmäßig die in den Einstellungen hinterlegten Vermieterdaten.
    """
    from database.settings import get_all_settings
    settings = get_all_settings()

    is_private = settings.get("landlord_type", "private") == "private"
    default_sender = "Kirdar Karadag" if is_private else "Hausverwaltung Sonnenhang"
    sender_name = sender_name or settings.get("landlord_name", default_sender)
    if is_private and sender_name == "Hausverwaltung Sonnenhang":
        sender_name = settings.get("landlord_contact_person") or "Kirdar Karadag"

    property_name = property_name or settings.get("landlord_name", "Wohnanlage Sonnenhang")
    sender_street = sender_street or settings.get("landlord_street", "Lindenstraße 14")
    sender_city = sender_city or f"{settings.get('landlord_zip', '10115')} {settings.get('landlord_city', 'Berlin')}"
    sender_phone = sender_phone or settings.get("landlord_phone", "+49 30 12345678")
    sender_email = sender_email or settings.get("landlord_email", "kirdar@karadag.de" if is_private else "verwaltung@sonnenhang.de")
    closing_phrase = settings.get("default_closing_phrase", "Mit freundlichen Grüßen")
    sign_title = sign_person or settings.get("default_signature_title")
    if not sign_title or (is_private and "hausverwaltung" in sign_title.lower()):
        sign_title = f"{sender_name}\n(Vermieter)" if is_private else sender_name
    os.makedirs(REPORTS_DIR, exist_ok=True)
    clean_first_line = recipient_name.split("\n")[0].strip()
    clean_name = "".join(c for c in clean_first_line if c.isalnum() or c in (" ", "_", "-")).strip().replace(" ", "_")[:35]
    today_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"Brief_{clean_name}_{today_str}.docx"
    docx_path = os.path.join(REPORTS_DIR, filename)

    doc = Document()
    
    # 1. DIN 5008 Form B Seitenränder: Oben 4.5 cm, Links 2.5 cm, Unten 2.5 cm, Rechts 2.0 cm
    for section in doc.sections:
        section.top_margin = Cm(4.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.0)

    # Betreff und Fließtext bereinigen
    clean_subject = re.sub(r"^(?:Betreff|Subject)\s*:\s*", "", subject, flags=re.IGNORECASE).strip()
    clean_text = re.sub(r"\n*(?:Mit freundlichen Grüßen|Freundliche Grüße|Herzliche Grüße)[,\s]*.*$", "", letter_text, flags=re.IGNORECASE | re.DOTALL).strip()

    # 2. Anschriftfeld (links, max. 8,5 cm) und Info-Block mit Datum (rechts, 8,0 cm)
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(8.5)
    table.columns[1].width = Cm(8.0)

    cell_left = table.cell(0, 0)
    cell_right = table.cell(0, 1)

    # Linke Spalte (Anschriftzone):
    # Rücksendeangabe (genau 8 pt unterstrichen)
    p_ret = cell_left.paragraphs[0]
    p_ret.paragraph_format.space_before = Pt(0)
    p_ret.paragraph_format.space_after = Pt(8)
    r_ret = p_ret.add_run(f"{sender_name} · {sender_street} · {sender_city}")
    r_ret.font.name = "Arial"
    r_ret.font.size = Pt(8.0)
    r_ret.font.underline = True
    r_ret.font.color.rgb = RGBColor(100, 116, 139)

    # Empfänger-Adresse (11 pt)
    p_rec = cell_left.add_paragraph()
    p_rec.paragraph_format.space_before = Pt(0)
    p_rec.paragraph_format.space_after = Pt(0)
    p_rec.paragraph_format.line_spacing = 1.0
    r_name = p_rec.add_run(f"{recipient_name}\n")
    r_name.font.name = "Arial"
    r_name.font.size = Pt(11.0)
    r_name.font.bold = True

    clean_addr = address_line.replace("<br/>", "\n").replace("<br>", "\n")
    addr_lines = [line.strip() for line in clean_addr.split("\n") if line.strip()]
    if not addr_lines:
        addr_lines = ["Musterstraße 1", "10115 Berlin"]

    for idx, line in enumerate(addr_lines):
        suffix = "\n" if idx < len(addr_lines) - 1 else ""
        r_line = p_rec.add_run(f"{line}{suffix}")
        r_line.font.name = "Arial"
        r_line.font.size = Pt(11.0)

    # Rechte Spalte (Info-Block & Datum, rechtsbündig):
    date_display = datetime.now().strftime("%d.%m.%Y")
    p_info = cell_right.paragraphs[0]
    p_info.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_info.paragraph_format.space_before = Pt(0)
    p_info.paragraph_format.space_after = Pt(0)
    p_info.paragraph_format.line_spacing = 1.0

    sender_title_txt = f"{sender_name} (Vermieter)\n" if is_private else f"{sender_name}\n"
    r_snd1 = p_info.add_run(sender_title_txt)
    r_snd1.font.name = "Arial"
    r_snd1.font.size = Pt(10.0)
    r_snd1.font.bold = True

    sender_city_name = ""
    if sender_city:
        parts = sender_city.strip().split()
        if len(parts) > 1 and parts[0].isdigit():
            sender_city_name = " ".join(parts[1:])
        elif parts:
            sender_city_name = parts[-1]
    effective_date_city = (settings.get("letter_city") or settings.get("landlord_city") or sender_city_name or "Berlin").strip()
    date_line_text = f"{effective_date_city}, den {date_display}"

    prop_txt = f"{property_name}\n" if (property_name and property_name != sender_name and not is_private) else ""
    info_text = (
        f"{prop_txt}"
        f"{sender_street}\n"
        f"{sender_city}\n"
        f"Tel.: {sender_phone}\n"
        f"E-Mail: {sender_email}\n\n"
        f"{date_line_text}"
    )
    r_snd2 = p_info.add_run(info_text)
    r_snd2.font.name = "Arial"
    r_snd2.font.size = Pt(9.5)
    r_snd2.font.color.rgb = RGBColor(51, 65, 85)

    # 3. Betreff: 2 Zeilen unter dem Datum, fett, ohne das Wort 'Betreff'
    p_subj = doc.add_paragraph()
    p_subj.paragraph_format.space_before = Pt(24)
    p_subj.paragraph_format.space_after = Pt(24)
    r_subj = p_subj.add_run(clean_subject)
    r_subj.font.name = "Arial"
    r_subj.bold = True
    r_subj.font.size = Pt(12.0)
    r_subj.font.color.rgb = RGBColor(15, 23, 42)

    # 4. Fließtext: 11 pt, einfacher Zeilenabstand (1.0), keine Erstzeileneinrückung, 1 Leerzeile nach Absatz
    paragraphs = [p.strip() for p in clean_text.split("\n\n") if p.strip()]
    if not paragraphs:
        paragraphs = [clean_text]

    for p in paragraphs:
        p_body = doc.add_paragraph()
        p_body.paragraph_format.left_indent = Pt(0)
        p_body.paragraph_format.first_line_indent = Pt(0)
        p_body.paragraph_format.line_spacing = 1.0
        p_body.paragraph_format.space_after = Pt(11)  # Leerzeile
        r_body = p_body.add_run(p)
        r_body.font.name = "Arial"
        r_body.font.size = Pt(11.0)
        r_body.font.color.rgb = RGBColor(15, 23, 42)

    # 5. Grußformel: nach einer Leerzeile, OHNE Komma, 3 Leerzeilen für handschriftliche Unterschrift
    p_gruss = doc.add_paragraph()
    p_gruss.paragraph_format.space_before = Pt(11)
    p_gruss.paragraph_format.space_after = Pt(33)  # 3 Leerzeilen für Unterschrift
    r_gruss = p_gruss.add_run(closing_phrase)
    r_gruss.font.name = "Arial"
    r_gruss.font.size = Pt(11.0)

    # Name des Unterzeichners / Absenders in Druckschrift
    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.space_before = Pt(0)
    p_sign.paragraph_format.space_after = Pt(0)
    r_sign = p_sign.add_run(sign_title)
    r_sign.font.name = "Arial"
    r_sign.font.size = Pt(11.0)
    r_sign.font.bold = True

    # 6. Optionale 2. Seite: Bestätigungsschreiben / Zustimmungserklärung
    if confirmation_data and confirmation_data.get("enabled", True):
        _build_confirmation_docx(
            doc=doc,
            confirmation_data=confirmation_data,
            sender_name=sender_name,
            sender_street=sender_street,
            sender_city=sender_city,
            is_private=is_private
        )

    doc.save(docx_path)
    return docx_path
